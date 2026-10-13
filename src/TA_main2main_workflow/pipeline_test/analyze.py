"""Pipeline test: AI log analysis + Excel report generation.

The AI (pipeline_test_analyze mode) reads the raw per-repo logs and
writes ``analysis.json`` into the step dir; the deterministic parser
(logs.py) is the fallback.  Both produce the same schema, consumed by
:func:`generate_excel_reports` to build the two Excel tables:

  - 错误信息总览表.xlsx — per-repo pass/fail counts
  - 错误信息汇总表.xlsx — per-failing-case error details + category
"""

from __future__ import annotations

import json
from pathlib import Path

from TA_main2main_workflow.agent.opencode_adapter import run_opencode_adapter
from TA_main2main_workflow.pipeline_test.logs import collect_repo_logs
from TA_main2main_workflow.utils import (
    FIX_LOG_DIR,
    PIPELINE_TEST_ANALYSIS_DIR_NAME,
    PIPELINE_TEST_ANALYSIS_FILE,
    PIPELINE_TEST_CATEGORY_LABELS,
    PIPELINE_TEST_DETAIL_XLSX,
    PIPELINE_TEST_OVERVIEW_XLSX,
    WORKSPACE_DIR,
)
from TA_main2main_workflow.utils.config import TAConfig
from TA_main2main_workflow.utils.context import WorkflowContext
from TA_main2main_workflow.utils.logging import get_logger

log = get_logger(__name__)
_REF = str(Path(__file__).parent.parent / "reference")

_CATEGORIES = set(PIPELINE_TEST_CATEGORY_LABELS)

_OVERVIEW_HEADERS = [
    "测试仓名称",
    "总用例数",
    "成功用例数",
    "失败用例数",
    "错误用例数",
    "跳过用例数",
    "失败率",
]
_DETAIL_HEADERS = ["测试仓名称", "失败的测试case", "报错信息", "问题分类"]
_COUNT_FIELDS = ("total", "passed", "failed", "errors", "skipped")


# ═══════════════════════════════════════════════════════════════════════════
# Schema validation
# ═══════════════════════════════════════════════════════════════════════════


def validate_analysis(data: object) -> tuple[bool, str]:
    """Validate *data* against the analysis.json schema.

    Returns ``(is_valid, reason)``.  Strict enough to trust the Excel
    generator, lenient enough to accept minor AI formatting drift.
    """
    if not isinstance(data, dict):
        return False, "analysis.json root must be an object"
    repos = data.get("repos")
    if not isinstance(repos, list) or not repos:
        return False, "'repos' must be a non-empty list"
    for repo in repos:
        if not isinstance(repo, dict):
            return False, "each repo must be an object"
        if not isinstance(repo.get("repo"), str) or not repo["repo"].strip():
            return False, "repo name missing"
        for field in _COUNT_FIELDS:
            if field not in repo:
                return False, f"repo {repo['repo']!r} missing count field '{field}'"
            if not isinstance(repo[field], int) or repo[field] < 0:
                return False, (
                    f"repo {repo['repo']!r} field '{field}' not a "
                    "non-negative int"
                )
        cases = repo.get("cases", [])
        if not isinstance(cases, list):
            return False, f"repo {repo['repo']!r} 'cases' must be a list"
        for case in cases:
            if not isinstance(case, dict):
                return False, f"repo {repo['repo']!r} case must be an object"
            if not isinstance(case.get("case"), str) or not case["case"].strip():
                return False, f"repo {repo['repo']!r} case id missing"
            if case.get("category") not in _CATEGORIES:
                return False, (
                    f"case {case.get('case')!r} category must be one of "
                    f"{sorted(_CATEGORIES)}"
                )
    return True, "valid"


# ═══════════════════════════════════════════════════════════════════════════
# AI analysis
# ═══════════════════════════════════════════════════════════════════════════


def run_ai_analysis(
    ctx: WorkflowContext,
    config: TAConfig,
    test_workspace: Path,
    step_dir: Path,
    deterministic: dict,
) -> dict:
    """Invoke the AI to analyze the raw logs; fall back to *deterministic*.

    The AI writes ``analysis.json`` into *step_dir*; it is validated
    against the schema before use.  Any failure (skip flag, adapter
    error, missing/invalid JSON) falls back to the deterministic result.
    """
    if config.skip_ai_analysis:
        log.info("SKIP_AI_ANALYSIS=true — using deterministic analysis")
        return deterministic

    repo_logs = collect_repo_logs(test_workspace)
    log_files = [str(p) for paths in repo_logs.values() for p in paths]
    if not log_files:
        log.warning("No test logs found — analysis skipped")
        return deterministic

    fix_dir = WORKSPACE_DIR / FIX_LOG_DIR / "pipeline-test-analyze"
    fix_dir.mkdir(parents=True, exist_ok=True)
    try:
        run_opencode_adapter(
            {
                "step_id": "pipeline-test-analyze",
                "previous_step_id": "",
                "previous_step_summary_path": "",
                "is_last_step": "true",
                "step_dir": str(step_dir),
                "fix_dir": str(fix_dir),
                "conflict_dir": "",
                "ascend_path": ctx.triton_ascend_path,
                "triton_path": ctx.triton_ascend_path,
                "reference_dir": _REF,
                "mode": "pipeline_test_analyze",
                "error_logs": json.dumps(
                    {"log_files": log_files, "hints": deterministic},
                    ensure_ascii=False,
                ),
                "target_commit": ctx.target_commit,
                "step_index": "1/1",
                "test_workspace": str(test_workspace),
                "ascend_npu_ir_fix": "false",
                "ascend_npu_ir_compat_ref": "",
            }
        )
    except Exception as exc:
        log.error(f"AI log analysis failed: {exc} — using deterministic fallback")
        return deterministic

    analysis_file = step_dir / PIPELINE_TEST_ANALYSIS_FILE
    if not analysis_file.exists():
        log.warning("AI produced no analysis.json — using deterministic fallback")
        return deterministic
    try:
        data = json.loads(analysis_file.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        log.warning(f"analysis.json unreadable ({exc}) — using deterministic fallback")
        return deterministic
    ok, reason = validate_analysis(data)
    if not ok:
        log.warning(f"analysis.json invalid ({reason}) — using deterministic fallback")
        return deterministic
    log.status(True, f"AI analysis valid: {len(data['repos'])} repo(s)")
    return data


# ═══════════════════════════════════════════════════════════════════════════
# Excel generation
# ═══════════════════════════════════════════════════════════════════════════


def generate_excel_reports(analysis: dict, out_dir: Path) -> tuple[Path, Path]:
    """Write the two Excel reports into *out_dir*.

    Returns ``(overview_path, detail_path)``.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    out_dir.mkdir(parents=True, exist_ok=True)
    overview_path = out_dir / PIPELINE_TEST_OVERVIEW_XLSX
    detail_path = out_dir / PIPELINE_TEST_DETAIL_XLSX

    header_font = Font(bold=True)

    # ── 错误信息总览表 ──────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "错误信息总览"
    ws.append(_OVERVIEW_HEADERS)
    for repo in analysis.get("repos", []):
        failed = repo.get("failed", 0)
        errors = repo.get("errors", 0)
        total = repo.get("total", 0)
        rate = f"{(failed + errors) / total * 100:.1f}%" if total else "0.0%"
        ws.append(
            [
                repo["repo"],
                total,
                repo.get("passed", 0),
                failed,
                errors,
                repo.get("skipped", 0),
                rate,
            ]
        )
    _style_sheet(ws, _OVERVIEW_HEADERS, header_font)
    wb.save(overview_path)

    # ── 错误信息汇总表 ──────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "错误信息汇总"
    ws.append(_DETAIL_HEADERS)
    for repo in analysis.get("repos", []):
        for case in repo.get("cases", []):
            ws.append(
                [
                    repo["repo"],
                    case.get("case", ""),
                    case.get("error", ""),
                    PIPELINE_TEST_CATEGORY_LABELS.get(
                        case.get("category", ""), case.get("category", "")
                    ),
                ]
            )
    _style_sheet(ws, _DETAIL_HEADERS, header_font)
    ws.column_dimensions[get_column_letter(2)].width = 60
    ws.column_dimensions[get_column_letter(3)].width = 80
    wb.save(detail_path)

    log.status(
        True,
        f"Excel reports: {overview_path.name}, {detail_path.name}",
    )
    return overview_path, detail_path


def _style_sheet(ws, headers: list[str], header_font) -> None:
    """Bold header, sane column widths, frozen header row."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    for cell in ws[1]:
        cell.font = header_font
    for idx, header in enumerate(headers, start=1):
        width = min(max(len(header) * 2 + 4, 12), 80)
        ws.column_dimensions[get_column_letter(idx)].width = width
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def analysis_dir(test_workspace: Path) -> Path:
    """The directory where the Excel reports live."""
    return test_workspace / PIPELINE_TEST_ANALYSIS_DIR_NAME
