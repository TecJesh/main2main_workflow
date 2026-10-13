"""Tests for the pipeline-test flow's pure helpers (ta-pipeline-test).

Covers the deterministic pieces that don't need an AI backend or a real
triton-ascend checkout:

  1. pytest summary-line parsing (with and without a summary)
  2. FAILED / ERROR short-summary parsing + category inference
  3. fallback_analysis over a fake test_workspace layout
  4. analysis.json schema validation
  5. Excel report generation (headers, rows, 失败率)

Run:  python3 -m pytest test/test_pipeline_test.py -v
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

# Ensure the src directory is importable (matches the other test files).
_SCRIPT_DIR = Path(__file__).resolve().parent
_PROJECT_ROOT = _SCRIPT_DIR.parent
_SRC = _PROJECT_ROOT / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from openpyxl import load_workbook

from TA_main2main_workflow.pipeline_test import analyze, logs
from TA_main2main_workflow.utils import (
    PIPELINE_TEST_CATEGORY_NPU_IR,
    PIPELINE_TEST_CATEGORY_OPERATOR,
    PIPELINE_TEST_CATEGORY_TRITON_ASCEND,
)

_SAMPLE_LOG = """\
tests/fla/test_a.py::test_pass PASSED
____ tests/fla/test_fused_cross_entropy_gen.py::test_cross_entropy_bwd[r7328_c512_ls0.0_bfloat16] ____
E   AssertionError: expected 4, got 2
E   assert 4 == 2
=========================== short test summary info ===========================
FAILED tests/fla/test_fused_cross_entropy_gen.py::test_cross_entropy_bwd[r7328_c512_ls0.0_bfloat16] - AssertionError: expected 4, got 2
ERROR tests/fla/test_b.py::test_err - npuir compile failed
======== 98 passed, 1 failed, 1 error, 2 skipped in 123.45s ========
"""


# ═══════════════════════════════════════════════════════════════════════════
# Summary + failure parsing
# ═══════════════════════════════════════════════════════════════════════════


def test_parse_pytest_summary():
    s = logs.parse_pytest_summary(_SAMPLE_LOG)
    assert s == {"total": 102, "passed": 98, "failed": 1, "errors": 1, "skipped": 2}


def test_parse_pytest_summary_missing():
    # -x early exit: no summary line → None, not a crash.
    assert logs.parse_pytest_summary("collected 10 items\n1 passed in 1.0s") is None


def test_parse_failed_cases_excerpts_and_kinds():
    cases = logs.parse_failed_cases(_SAMPLE_LOG)
    assert len(cases) == 2
    failed = next(c for c in cases if c["kind"] == "failed")
    error = next(c for c in cases if c["kind"] == "error")
    assert failed["case"].startswith("tests/fla/test_fused_cross_entropy_gen.py::")
    assert "expected 4" in failed["error"]
    assert error["case"] == "tests/fla/test_b.py::test_err"
    assert "npuir" in error["error"]


def test_infer_category():
    assert logs.infer_category("npuir compile failed") == PIPELINE_TEST_CATEGORY_NPU_IR
    assert (
        logs.infer_category("boom in lib/Conversion/TritonGPUToLLVM.cpp")
        == PIPELINE_TEST_CATEGORY_OPERATOR
    )
    assert (
        logs.infer_category("AssertionError: expected 4, got 2")
        == PIPELINE_TEST_CATEGORY_TRITON_ASCEND
    )


# ═══════════════════════════════════════════════════════════════════════════
# fallback_analysis
# ═══════════════════════════════════════════════════════════════════════════


def _make_workspace(log_text: str, extra_files: dict[str, str] | None = None) -> Path:
    tw = Path(tempfile.mkdtemp())
    repo_log = tw / "test_log" / "Q3TritonKernel_test_log"
    repo_log.mkdir(parents=True)
    (repo_log / "test_fla.log").write_text(log_text, encoding="utf-8")
    for name, content in (extra_files or {}).items():
        (repo_log / name).write_text(content, encoding="utf-8")
    return tw


def test_fallback_analysis_counts_and_repo_name():
    tw = _make_workspace(
        _SAMPLE_LOG, {"rerun_case-0.log": "2 passed in 1.0s"}
    )
    analysis = logs.fallback_analysis(tw)
    assert len(analysis["repos"]) == 1
    repo = analysis["repos"][0]
    # Repo name = dir name minus the _test_log suffix; rerun logs excluded.
    assert repo["repo"] == "Q3TritonKernel"
    assert repo["total"] == 102
    assert repo["skipped"] == 2
    assert len(repo["cases"]) == 2
    assert logs.has_failures(analysis)


def test_fallback_analysis_no_summary_counts_from_failed_lines():
    text = (
        "FAILED tests/x.py::test_x - boom\n"
        "ERROR tests/x.py::test_y - npuir error\n"
    )
    analysis = logs.fallback_analysis(_make_workspace(text))
    repo = analysis["repos"][0]
    assert repo["total"] == 2
    assert repo["failed"] == 1
    assert repo["errors"] == 1
    assert repo["cases"][1]["category"] == PIPELINE_TEST_CATEGORY_NPU_IR


# ═══════════════════════════════════════════════════════════════════════════
# Schema validation
# ═══════════════════════════════════════════════════════════════════════════


def _valid_analysis() -> dict:
    tw = _make_workspace(_SAMPLE_LOG)
    return logs.fallback_analysis(tw)


def test_validate_analysis_accepts_fallback():
    ok, reason = analyze.validate_analysis(_valid_analysis())
    assert ok, reason


def test_validate_analysis_rejects_bad_category():
    bad = {
        "repos": [
            {
                "repo": "X",
                "total": 1, "passed": 1, "failed": 0, "errors": 0, "skipped": 0,
                "cases": [{"case": "a::b", "error": "", "category": "weird"}],
            }
        ]
    }
    ok, reason = analyze.validate_analysis(bad)
    assert not ok
    assert "category" in reason


def test_validate_analysis_rejects_missing_counts():
    ok, reason = analyze.validate_analysis({"repos": [{"repo": "X"}]})
    assert not ok
    assert "count" in reason


# ═══════════════════════════════════════════════════════════════════════════
# Excel generation
# ═══════════════════════════════════════════════════════════════════════════


def test_generate_excel_reports():
    tw = Path(tempfile.mkdtemp())
    out_dir = tw / "analysis"
    overview_path, detail_path = analyze.generate_excel_reports(
        _valid_analysis(), out_dir
    )
    assert overview_path.exists() and detail_path.exists()

    ws = load_workbook(overview_path).active
    headers = [c.value for c in ws[1]]
    assert headers == [
        "测试仓名称", "总用例数", "成功用例数", "失败用例数",
        "错误用例数", "跳过用例数", "失败率",
    ]
    row = [c.value for c in ws[2]]
    assert row[0] == "Q3TritonKernel"
    # (1 failed + 1 error) / 102 = 2.0%
    assert row[6] == "2.0%"

    ws2 = load_workbook(detail_path).active
    headers2 = [c.value for c in ws2[1]]
    assert headers2 == ["测试仓名称", "失败的测试case", "报错信息", "问题分类"]
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    categories = {r[3] for r in rows}
    assert categories == {"Triton—Ascend问题", "NPU-IR问题"}


def test_generate_excel_reports_all_pass_zero_rate():
    tw = _make_workspace("10 passed in 1.0s")
    out_dir = tw / "analysis"
    overview_path, detail_path = analyze.generate_excel_reports(
        logs.fallback_analysis(tw), out_dir
    )
    ws = load_workbook(overview_path).active
    row = [c.value for c in ws[2]]
    assert row[6] == "0.0%"
    # No failing cases → detail sheet has only the header row.
    ws2 = load_workbook(detail_path).active
    assert ws2.max_row == 1
